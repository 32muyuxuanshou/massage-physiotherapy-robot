"""Build the SMPL-X/SKEL illustrated tutorial for first-time doctor annotators."""

from pathlib import Path

from docx import Document
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK
from docx.enum.section import WD_SECTION_START
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


ROOT = Path(__file__).resolve().parent
IMAGES = ROOT / "screenshots" / "prepared"
OUTPUT = ROOT / "医生版_SMPL-X与SKEL穴位标注图文教程_v1.2.docx"

# compact_reference_guide preset, resolved explicitly.
PAGE_W, PAGE_H = Inches(8.5), Inches(11)
MARGIN = Inches(1)
CONTENT_DXA = 9360
TABLE_INDENT_DXA = 120
CELL_MARGINS = {"top": 80, "bottom": 80, "start": 120, "end": 120}
NAVY = "0B2545"
BLUE = "2E74B5"
DARK_BLUE = "1F4D78"
MUTED = "5D6876"
FILL = "F4F6F9"
TABLE_FILL = "E8EEF5"
GOLD = "7A5A00"
RED = "9B1C1C"
GREEN = "1F5D4E"
WHITE = "FFFFFF"


def set_run(run, size=None, color=NAVY, bold=None, italic=None, font="Calibri"):
    run.font.name = font
    fonts = run._element.get_or_add_rPr().rFonts
    fonts.set(qn("w:ascii"), font)
    fonts.set(qn("w:hAnsi"), font)
    fonts.set(qn("w:eastAsia"), "Microsoft YaHei")
    if size is not None:
        run.font.size = Pt(size)
    if color:
        run.font.color.rgb = RGBColor.from_string(color)
    if bold is not None:
        run.bold = bold
    if italic is not None:
        run.italic = italic


def set_cell_margins(cell, **kwargs):
    tc = cell._tc
    tc_pr = tc.get_or_add_tcPr()
    tc_mar = tc_pr.first_child_found_in("w:tcMar")
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for edge in ("top", "start", "bottom", "end"):
        value = kwargs.get(edge)
        if value is None:
            continue
        node = tc_mar.find(qn(f"w:{edge}"))
        if node is None:
            node = OxmlElement(f"w:{edge}")
            tc_mar.append(node)
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")


def set_cell_width(cell, dxa):
    tc_pr = cell._tc.get_or_add_tcPr()
    tc_w = tc_pr.find(qn("w:tcW"))
    if tc_w is None:
        tc_w = OxmlElement("w:tcW")
        tc_pr.append(tc_w)
    tc_w.set(qn("w:w"), str(dxa))
    tc_w.set(qn("w:type"), "dxa")


def shade_cell(cell, fill):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), fill)


def set_table_geometry(table, widths):
    table.alignment = WD_TABLE_ALIGNMENT.LEFT
    table.autofit = False
    tbl_pr = table._tbl.tblPr
    tbl_w = tbl_pr.find(qn("w:tblW"))
    if tbl_w is None:
        tbl_w = OxmlElement("w:tblW")
        tbl_pr.append(tbl_w)
    tbl_w.set(qn("w:w"), str(sum(widths)))
    tbl_w.set(qn("w:type"), "dxa")
    tbl_ind = tbl_pr.find(qn("w:tblInd"))
    if tbl_ind is None:
        tbl_ind = OxmlElement("w:tblInd")
        tbl_pr.append(tbl_ind)
    tbl_ind.set(qn("w:w"), str(TABLE_INDENT_DXA))
    tbl_ind.set(qn("w:type"), "dxa")
    grid = table._tbl.tblGrid
    for child in list(grid):
        grid.remove(child)
    for width in widths:
        col = OxmlElement("w:gridCol")
        col.set(qn("w:w"), str(width))
        grid.append(col)
    for row in table.rows:
        for cell, width in zip(row.cells, widths):
            set_cell_width(cell, width)
            set_cell_margins(cell, **CELL_MARGINS)
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER


def set_repeat_table_header(row):
    tr_pr = row._tr.get_or_add_trPr()
    repeat = OxmlElement("w:tblHeader")
    repeat.set(qn("w:val"), "true")
    tr_pr.append(repeat)


def add_page_field(paragraph):
    run = paragraph.add_run("第 ")
    set_run(run, 8.5, MUTED)
    begin = OxmlElement("w:fldChar")
    begin.set(qn("w:fldCharType"), "begin")
    instr = OxmlElement("w:instrText")
    instr.set(qn("xml:space"), "preserve")
    instr.text = " PAGE "
    separate = OxmlElement("w:fldChar")
    separate.set(qn("w:fldCharType"), "separate")
    text = OxmlElement("w:t")
    text.text = "1"
    end = OxmlElement("w:fldChar")
    end.set(qn("w:fldCharType"), "end")
    run._r.extend([begin, instr, separate, text, end])
    set_run(paragraph.add_run(" 页"), 8.5, MUTED)


def style_document(doc):
    section = doc.sections[0]
    section.page_width, section.page_height = PAGE_W, PAGE_H
    section.top_margin = section.right_margin = section.bottom_margin = section.left_margin = MARGIN
    section.header_distance = section.footer_distance = Inches(0.492)
    styles = doc.styles
    normal = styles["Normal"]
    normal.font.name = "Calibri"
    normal._element.rPr.rFonts.set(qn("w:eastAsia"), "Microsoft YaHei")
    normal.font.size = Pt(11)
    normal.paragraph_format.space_before = Pt(0)
    normal.paragraph_format.space_after = Pt(6)
    normal.paragraph_format.line_spacing = 1.25
    for name, size, color, before, after in (
        ("Heading 1", 16, BLUE, 18, 10),
        ("Heading 2", 13, BLUE, 14, 7),
        ("Heading 3", 12, DARK_BLUE, 10, 5),
    ):
        style = styles[name]
        style.font.name = "Calibri"
        style._element.rPr.rFonts.set(qn("w:eastAsia"), "Microsoft YaHei")
        style.font.size = Pt(size)
        style.font.bold = True
        style.font.color.rgb = RGBColor.from_string(color)
        style.paragraph_format.space_before = Pt(before)
        style.paragraph_format.space_after = Pt(after)
        style.paragraph_format.keep_with_next = True
        if name == "Heading 1":
            style.paragraph_format.page_break_before = True
    for name in ("List Bullet", "List Number"):
        style = styles[name]
        style.font.name = "Calibri"
        style._element.rPr.rFonts.set(qn("w:eastAsia"), "Microsoft YaHei")
        style.font.size = Pt(11)
        style.paragraph_format.left_indent = Inches(0.375)
        style.paragraph_format.first_line_indent = Inches(-0.188)
        style.paragraph_format.space_after = Pt(4)
        style.paragraph_format.line_spacing = 1.25
    caption = styles["Caption"]
    caption.font.name = "Calibri"
    caption._element.rPr.rFonts.set(qn("w:eastAsia"), "Microsoft YaHei")
    caption.font.size = Pt(9)
    caption.font.color.rgb = RGBColor.from_string(MUTED)
    caption.paragraph_format.space_before = Pt(3)
    caption.paragraph_format.space_after = Pt(9)
    caption.paragraph_format.line_spacing = 1.15
    header = section.header.paragraphs[0]
    header.alignment = WD_ALIGN_PARAGRAPH.LEFT
    header.paragraph_format.space_after = Pt(0)
    set_run(header.add_run("医生穴位标注操作手册  |  SMPL-X 与 SKEL"), 8.5, MUTED, True)
    p_pr = header._p.get_or_add_pPr()
    border = OxmlElement("w:pBdr")
    bottom = OxmlElement("w:bottom")
    bottom.set(qn("w:val"), "single")
    bottom.set(qn("w:sz"), "4")
    bottom.set(qn("w:space"), "5")
    bottom.set(qn("w:color"), "D7DBE2")
    border.append(bottom)
    p_pr.append(border)
    footer = section.footer.paragraphs[0]
    footer.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    footer.paragraph_format.space_before = Pt(0)
    add_page_field(footer)


def add_para(doc, text, lead=None, color=NAVY, after=6, italic=False):
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(after)
    if lead and text.startswith(lead):
        set_run(p.add_run(lead), color=color, bold=True)
        set_run(p.add_run(text[len(lead):]), color=color, italic=italic)
    else:
        set_run(p.add_run(text), color=color, italic=italic)
    return p


def add_callout(doc, label, text, kind="info"):
    colors = {"info": (FILL, BLUE), "warning": ("FFF4E5", RED), "success": ("EEF8F4", GREEN)}
    fill, stripe = colors[kind]
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(5)
    p.paragraph_format.space_after = Pt(9)
    p.paragraph_format.left_indent = Inches(0.12)
    p.paragraph_format.right_indent = Inches(0.12)
    p.paragraph_format.line_spacing = 1.2
    p_pr = p._p.get_or_add_pPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:fill"), fill)
    p_pr.append(shd)
    p_bdr = OxmlElement("w:pBdr")
    left = OxmlElement("w:left")
    left.set(qn("w:val"), "single")
    left.set(qn("w:sz"), "18")
    left.set(qn("w:space"), "8")
    left.set(qn("w:color"), stripe)
    p_bdr.append(left)
    p_pr.append(p_bdr)
    set_run(p.add_run(f"{label}："), 10.5, stripe, True)
    set_run(p.add_run(text), 10.5, NAVY)
    return p


def new_list_numbering(doc, numbered=False):
    numbering = doc.part.numbering_part.element
    abstract_ids = [int(n.get(qn("w:abstractNumId"))) for n in numbering.findall(qn("w:abstractNum"))]
    num_ids = [int(n.get(qn("w:numId"))) for n in numbering.findall(qn("w:num"))]
    abstract_id = max(abstract_ids, default=0) + 1
    num_id = max(num_ids, default=0) + 1
    abstract = OxmlElement("w:abstractNum")
    abstract.set(qn("w:abstractNumId"), str(abstract_id))
    multi = OxmlElement("w:multiLevelType")
    multi.set(qn("w:val"), "singleLevel")
    abstract.append(multi)
    level = OxmlElement("w:lvl")
    level.set(qn("w:ilvl"), "0")
    start = OxmlElement("w:start")
    start.set(qn("w:val"), "1")
    level.append(start)
    num_fmt = OxmlElement("w:numFmt")
    num_fmt.set(qn("w:val"), "decimal" if numbered else "bullet")
    level.append(num_fmt)
    level_text = OxmlElement("w:lvlText")
    level_text.set(qn("w:val"), "%1." if numbered else "•")
    level.append(level_text)
    suffix = OxmlElement("w:suff")
    suffix.set(qn("w:val"), "tab")
    level.append(suffix)
    level_jc = OxmlElement("w:lvlJc")
    level_jc.set(qn("w:val"), "left")
    level.append(level_jc)
    p_pr = OxmlElement("w:pPr")
    tabs = OxmlElement("w:tabs")
    tab = OxmlElement("w:tab")
    tab.set(qn("w:val"), "num")
    tab.set(qn("w:pos"), "540")
    tabs.append(tab)
    p_pr.append(tabs)
    indent = OxmlElement("w:ind")
    indent.set(qn("w:left"), "540")
    indent.set(qn("w:hanging"), "270")
    p_pr.append(indent)
    spacing = OxmlElement("w:spacing")
    spacing.set(qn("w:after"), "80")
    spacing.set(qn("w:line"), "300")
    spacing.set(qn("w:lineRule"), "auto")
    p_pr.append(spacing)
    level.append(p_pr)
    abstract.append(level)
    numbering.append(abstract)
    num = OxmlElement("w:num")
    num.set(qn("w:numId"), str(num_id))
    abstract_ref = OxmlElement("w:abstractNumId")
    abstract_ref.set(qn("w:val"), str(abstract_id))
    num.append(abstract_ref)
    numbering.append(num)
    return num_id


def add_bullets(doc, items, numbered=False):
    num_id = new_list_numbering(doc, numbered=numbered)
    for item in items:
        p = doc.add_paragraph()
        p.paragraph_format.space_after = Pt(4)
        p.paragraph_format.line_spacing = 1.25
        p_pr = p._p.get_or_add_pPr()
        num_pr = OxmlElement("w:numPr")
        ilvl = OxmlElement("w:ilvl")
        ilvl.set(qn("w:val"), "0")
        num_id_el = OxmlElement("w:numId")
        num_id_el.set(qn("w:val"), str(num_id))
        num_pr.extend([ilvl, num_id_el])
        p_pr.append(num_pr)
        set_run(p.add_run(item))


def add_table(doc, headers, rows, widths=None):
    if widths is None:
        widths = [CONTENT_DXA // len(headers)] * len(headers)
        widths[-1] += CONTENT_DXA - sum(widths)
    table = doc.add_table(rows=1, cols=len(headers))
    table.style = "Table Grid"
    set_table_geometry(table, widths)
    hdr = table.rows[0]
    set_repeat_table_header(hdr)
    for cell, header in zip(hdr.cells, headers):
        shade_cell(cell, TABLE_FILL)
        p = cell.paragraphs[0]
        p.paragraph_format.space_after = Pt(0)
        set_run(p.add_run(header), 10.2, NAVY, True)
    for row in rows:
        cells = table.add_row().cells
        for cell, value in zip(cells, row):
            p = cell.paragraphs[0]
            p.paragraph_format.space_after = Pt(0)
            set_run(p.add_run(str(value)), 9.8, NAVY)
    set_table_geometry(table, widths)
    return table


def add_figure(doc, filename, caption, width=6.5):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_before = Pt(4)
    p.paragraph_format.space_after = Pt(0)
    run = p.add_run()
    run.add_picture(str(IMAGES / filename), width=Inches(width))
    cap = doc.add_paragraph(style="Caption")
    cap.alignment = WD_ALIGN_PARAGRAPH.CENTER
    set_run(cap.add_run(caption), 9, MUTED)


def page_break(doc):
    # Heading 1 has page_break_before.  A standalone page-break paragraph can
    # become an empty page when the preceding content already fills its page.
    return None


def checkbox_rows(items):
    return [("□", item) for item in items]


doc = Document()
style_document(doc)

# Cover — editorial_cover header pattern, while retaining compact_reference_guide tokens.
p = doc.add_paragraph()
p.alignment = WD_ALIGN_PARAGRAPH.CENTER
p.paragraph_format.space_before = Pt(80)
p.paragraph_format.space_after = Pt(18)
set_run(p.add_run("科研标注操作手册"), 11, GOLD, True)
p = doc.add_paragraph()
p.alignment = WD_ALIGN_PARAGRAPH.CENTER
p.paragraph_format.space_after = Pt(8)
set_run(p.add_run("SMPL-X 与 SKEL"), 30, NAVY, True)
p = doc.add_paragraph()
p.alignment = WD_ALIGN_PARAGRAPH.CENTER
p.paragraph_format.space_after = Pt(30)
set_run(p.add_run("医生穴位标注图文教程"), 17, BLUE, True)
add_callout(
    doc,
    "用途",
    "在本机 Blender 4.5.12 中，由专业医生在可变形人体表面记录穴位或按摩目标点。软件只记录医生判断，不自动给出临床位置。",
    "info",
)
p = doc.add_paragraph()
p.alignment = WD_ALIGN_PARAGRAPH.CENTER
p.paragraph_format.space_before = Pt(125)
p.paragraph_format.space_after = Pt(4)
set_run(p.add_run("版本 1.2  |  2026-08-13"), 11, NAVY, True)
p = doc.add_paragraph()
p.alignment = WD_ALIGN_PARAGRAPH.CENTER
set_run(p.add_run("适用环境：项目自带 Blender 4.5.12 LTS 与穴位标注插件 v0.3.2"), 9.5, MUTED, italic=True)

page_break(doc)
doc.add_heading("先看这一页：10 分钟完成一次标注", 1)
add_callout(
    doc,
    "第一条安全规则",
    "打开模板后先“文件 → 另存为”，保存工作副本；不要直接在名称含“空白模板”的文件上按 Ctrl+S。",
    "warning",
)
add_table(
    doc,
    ["步", "要做什么", "完成标志"],
    [
        ("1", "双击对应的一键启动文件，打开 SMPL-X 或 SKEL 模板", "Blender 窗口出现"),
        ("2", "立即另存为工作副本", "标题栏显示自己的新文件名"),
        ("3", "单击人体；按 Home 居中，按 N 打开右侧栏", "看到“全身穴位标注”"),
        ("4", "选择视角和模型；先调整体型、再调整姿态", "人体接近本次标注参考体型/姿态"),
        ("5", "SMPL-X 更新姿态修正；SKEL 执行皮肤或完整更新", "人体表面已刷新"),
        ("6", "填写编码、中文名称、身体分区、经脉/区域、侧别、备注", "字段完整"),
        ("7", "单击“在人体表面选点”，再单击皮肤", "出现小十字和医学名称"),
        ("8", "重复标 3 个点；切前/后/左/右逐点复核", "列表中有 3 项且侧别正确"),
        ("9", "Ctrl+S 保存 `.blend`，再导出穴位 JSON", "两个文件都存在"),
        ("10", "关闭并重新打开副本，抽查 3 点", "名称、位置、数量仍正确"),
    ],
    [500, 5700, 3160],
)
add_callout(doc, "建议主线", "全身、手部和头面部标注优先用 SMPL-X；需要骨骼/生物力学对照时再开 SKEL。", "success")

doc.add_heading("本教程的演示点", 2)
add_para(doc, "DEMO-01～DEMO-03 均为非医学演示点，只用于说明按钮、保存和复查流程，不代表任何标准穴位。")

page_break(doc)
doc.add_heading("1. 先选对模型", 1)
add_table(
    doc,
    ["需求", "选择", "原因与限制"],
    [
        ("全身穴位、背部按摩点", "SMPL-X（主模板）", "表面完整，可统一记录全身点位。"),
        ("手部精细点", "SMPL-X", "有独立手指关节；标点前需放大手部。"),
        ("头面部精细点", "SMPL-X", "支持下颌与面部表情参数；标点前需放大并切换视角。"),
        ("骨骼、关节、生物力学对照", "SKEL", "可显示内部骨骼和关节；计算更新比 SMPL-X 慢。"),
        ("SKEL 上的手指/表情", "不要用 SKEL", "SKEL 没有独立手指关节和面部表情；改用 SMPL-X。"),
    ],
    [2300, 1900, 5160],
)
doc.add_heading("1.1 现有模板与一键启动", 2)
add_para(doc, "本机已经部署好，医生不需要手工输入完整路径。打开 Windows 文件资源管理器，按图示依次进入“办公 (E:)”→“项目-按摩理疗机器人”→“AI感知模块”→“模型资源”，再进入 SMPL-X 或 SKEL 文件夹。")
add_figure(doc, "00_explorer_launch.jpg", "图 1　按黄色文件夹逐级进入，并在文件列表底部双击对应的 `.cmd` 一键启动文件。", width=6.5)
add_table(
    doc,
    ["模板", "双击这个文件"],
    [
        ("SMPL-X 中性", "模型资源/SMPL-X/打开_中性_SMPL-X全身标注.cmd"),
        ("SMPL-X 女性", "模型资源/SMPL-X/打开_女性_SMPL-X全身标注.cmd"),
        ("SMPL-X 男性", "模型资源/SMPL-X/打开_男性_SMPL-X全身标注.cmd"),
        ("SKEL 女性", "模型资源/SKEL/打开_女性_SKEL生物力学标注.cmd"),
        ("SKEL 男性", "模型资源/SKEL/打开_男性_SKEL生物力学标注.cmd"),
    ],
    [2200, 7160],
)
add_callout(doc, "模型性别", "性别来自你打开的模板，不是软件自动判断。要更换中性/女性/男性模型，请关闭当前工作副本，再打开对应模板；不要在同一场景中按“添加”生成第二个人体。", "warning")
add_para(doc, "推荐默认：如果项目尚未规定性别分层，先用 SMPL-X 中性模板建立主标注集；再由研究方案决定是否补做女性、男性差异。")

page_break(doc)
doc.add_heading("2. 建立工作副本与文件命名", 1)
doc.add_heading("2.1 打开后第一件事：另存为", 2)
add_bullets(
    doc,
    [
        "单击“文件 → 另存为”，或按 Ctrl+Shift+S。",
        "进入本项目指定的标注输出目录，不要保存到模板目录。",
        "输入新文件名并确认；检查 Blender 标题栏已经变成新文件名。",
        "之后可用 Ctrl+S 保存当前工作副本。",
    ],
    numbered=True,
)
add_callout(doc, "不要覆盖空白模板", "模板是所有医生的共同起点。若误覆盖，会把你的体型、姿态和标注带给下一位医生。", "warning")
add_callout(doc, "建议输出目录", "本机科研试运行可在图 1 所示的“AI感知模块”文件夹下新建“医生标注输出/标注者编号/日期/”。若课题负责人下发了加密盘或数据平台路径，以书面任务单为准；教程不替代项目数据管理制度。", "info")
doc.add_heading("2.2 推荐命名", 2)
add_table(
    doc,
    ["文件类型", "示例"],
    [
        ("Blender 工作副本", "SMPLX_中性_背部_标注者A_20260813_v01.blend"),
        ("穴位数据", "SMPLX_中性_背部_标注者A_20260813_v01.json"),
        ("SKEL 工作副本", "SKEL_女性_躯干_标注者A_20260813_v01.blend"),
    ],
    [2200, 7160],
)
add_bullets(doc, ["使用研究编号或标注者编号，不在文件名中写患者姓名。", "每次大改递增 v01、v02；不要用“最终版”“最终版2”。", "`.blend` 和 `.json` 使用相同主文件名，便于成对归档。"])

page_break(doc)
doc.add_heading("3. 所有模型共用的界面操作", 1)
add_figure(doc, "01_smplx_open_and_annotate.jpg", "图 2　按 Home 居中人体，按 N 打开“项目/Item”右侧栏，然后进入全身穴位标注。")
doc.add_heading("3.1 人体太小、面板不见了怎么办", 2)
add_table(
    doc,
    ["现象", "处理"],
    [
        ("人体只是一个很小的点", "单击人体，再把鼠标放在视图区，按 Home。"),
        ("右侧没有穴位标注", "鼠标放在视图区，按 N；点击竖排“项目/Item”。"),
        ("面板太窄、文字被截断", "拖动右侧栏左边界向左，放宽面板。"),
        ("滚轮缩放不了人体", "确认鼠标停在视图区，而不是右侧栏。"),
        ("模型转乱了", "在穴位面板顶部按“前/后/左/右”恢复标准视角。"),
    ],
    [2900, 6460],
)
doc.add_heading("3.2 手部/头面部的最小视图操作", 2)
add_table(
    doc,
    ["目的", "鼠标操作", "成功标志"],
    [
        ("缩放", "鼠标停在三维视图区，滚动滚轮", "手或面部变大/变小，右侧栏不滚动。"),
        ("旋转视角", "按住鼠标中键并拖动", "能从正面转到斜侧面检查。"),
        ("平移画面", "按住 Shift + 鼠标中键并拖动", "目标手/脸移动到视图中央。"),
        ("回到全身", "先选人体，再按 Home", "完整人体重新居中。"),
        ("定位已有点", "在“已标注”列表选中后按“定位”", "视图聚焦到该点附近。"),
    ],
    [1900, 3900, 3560],
)
add_callout(doc, "精细点复核", "放大到能分辨目标解剖表面后再落点；至少用正面与斜侧面各看一次。若迷失视角，按 Home 回到全身，再重新缩放。", "info")

page_break(doc)
doc.add_heading("4. SMPL-X：调整体型与姿态", 1)
add_figure(doc, "02_smplx_shape_pose.jpg", "图 3　SMPL-X 体型与姿态面板。输入身高体重后生成统计体型近似。", width=5.65)
doc.add_heading("4.1 体型", 2)
add_bullets(
    doc,
    [
        "切到右侧竖排“SMPL模型”标签。模板内已经有人体，不要按上方“添加”。",
        "在“形状”中输入目标身高（米）和目标体重（千克）。",
        "单击“根据身高体重生成体型”；观察正面、侧面和背面。",
        "不满意时单击“重置”，再重新输入。",
    ],
    numbered=True,
)
add_callout(doc, "研究边界", "身高体重生成的是统计近似体型，不是患者三维扫描，也不代表个体组织厚度。医生标注时需记录这是“模板参考位置”。", "warning")
doc.add_heading("4.2 姿态", 2)
add_callout(doc, "先切换面板", "图 3 所在的“SMPL模型”页签只负责体型、姿态修正和手部预设，本来就不显示关节角度滑条。调整完体型后，按 N 打开右侧栏，切回“项目/Item”页签，在“全身穴位标注”中向下滚动。", "warning")
add_figure(doc, "03_smplx_pose_sliders.jpg", "图 4　切回“项目/Item”页签，单击“显示姿态滑动条”后才会展开关节角度。", width=6.35)
add_bullets(doc, ["先完成大体型调整，再展开滑动条，调整躯干、肩、肘、腕和下颌等姿态；手指整体状态使用“SMPL模型”页签中的“手部姿态”预设。", "保持“启用姿态修正”；改完姿态后回到“SMPL模型”页签，单击“更新姿态修正”。", "尽量在落点前确定姿态；后改姿态时，回到“全身穴位标注”单击“刷新”并逐点复核。"])
doc.add_heading("4.3 姿态滑条试调规则", 2)
add_table(
    doc,
    ["次序", "动作", "判断"],
    [
        ("1", "保存一个可回退版本", "标题栏是工作副本，Ctrl+S 已完成。"),
        ("2", "一次只改一个关节/轴滑条", "能明确是哪一个参数导致变化。"),
        ("3", "先小幅正值，再试负值", "找到屏幕上符合预期的方向。"),
        ("4", "方向或外观不对就恢复 0", "不把错误姿态带入后续标点。"),
        ("5", "更新姿态修正并看正、侧、背", "表面稳定、无明显扭曲后再落点。"),
    ],
    [900, 4200, 4260],
)
page_break(doc)
doc.add_heading("5. SMPL-X：填写信息并在表面落点", 1)
doc.add_heading("5.1 一个点的完整操作", 2)
add_bullets(
    doc,
    [
        "在“全身穴位标注 → 1. 选择人体”中确认目标是当前 SMPL-X 人体。",
        "填写标准编码：采用课题组统一编码；不要留空后再靠肉眼猜。",
        "填写中文名称：这里填医生认可的医学名称。人体旁将直接显示该名称，不再显示随机对象名。",
        "选择身体分区；填写经脉/区域；选择左侧、右侧或中线；在医生备注中写定位依据或疑问。",
        "单击“在人体表面选点”。鼠标变为准星后，在目标皮肤表面单击一次。",
        "看到小十字和名称后，立刻切一个相邻视角，确认点没有落到身体背面或另一层表面。",
    ],
    numbered=True,
)
add_callout(doc, "小十字不是球体", "当前版本用小型十字标记，默认大小 0.005 m。标记只是可视化符号，真正数据记录在皮肤三角面及其面内位置上。", "info")
add_callout(doc, "字段数据字典", "“标准编码、身体分区、经脉/区域”的正式写法必须由课题负责人在标注任务单中下发。没有数据字典时只做本教程 DEMO 练习，不开始正式医生标注，也不要自创缩写。", "warning")
doc.add_heading("5.2 左右侧别", 2)
add_para(doc, "侧别以人体自身左右为准，不以屏幕左右为准。正面看人体时，人体右侧通常出现在屏幕左边。软件不会自动镜像；左右同名穴位必须分别落点、分别检查。")
doc.add_heading("5.3 点选模式如何取消", 2)
add_bullets(doc, ["鼠标变成准星后，按 Esc 或单击鼠标右键可取消，不会创建新点。", "若单击空白处，软件会提示“没有命中人体”；仍处于准星模式，可继续单击皮肤，或按 Esc 退出。", "一次成功落点后准星自动结束。需要下一个点时重新单击“在人体表面选点”。"])

page_break(doc)
doc.add_heading("6. SMPL-X：3 点演示与复核", 1)
add_figure(doc, "03_smplx_three_points_front.jpg", "图 5　正面 3 个非医学演示点：头面、躯干和手部。", width=5.45)
add_figure(doc, "04_smplx_three_points_back.jpg", "图 6　切到背面复核：名称、位置、身体分区与侧别应一致。", width=5.45)
add_callout(doc, "改错方法", "元数据写错：在“已标注”列表选中该点，修改下方信息。位置点错：删除该点后重新落点；当前版本不支持直接拖动点位。", "warning")
doc.add_heading("6.1 练习用 3 点字段表", 2)
add_table(
    doc,
    ["编码", "中文名称", "身体分区", "经脉/区域", "侧别", "落点要求"],
    [
        ("DEMO-01", "演示点一（非医学）", "头面", "教程演示区域", "中线", "额部任一清晰表面；不按医学穴位定位。"),
        ("DEMO-02", "演示点二（非医学）", "躯干", "教程演示区域", "左侧", "左侧腹部任一清晰表面。"),
        ("DEMO-03", "演示点三（非医学）", "右手", "教程演示区域", "右侧", "右手背任一清晰表面。"),
    ],
    [1250, 2100, 1150, 1700, 900, 2260],
)
add_callout(doc, "练习目标", "只要求字段、侧别和三类表面正确，不要求与截图落在同一三角面。正式标注必须改用课题组下发的医学数据字典。", "info")

page_break(doc)
doc.add_heading("7. 标注列表、显示与数据交换", 1)
doc.add_heading("7.1 管理已有点", 2)
add_table(
    doc,
    ["按钮/区域", "作用", "注意"],
    [
        ("已标注列表", "显示所有点，并显示分区与侧别", "点名称来自“中文名称”字段。"),
        ("定位", "把视图对准所选点", "用于逐点抽查。"),
        ("删除", "删除所选点", "位置点错时使用；删除前确认编码。"),
        ("刷新", "根据当前人体表面刷新标记", "体型/姿态改动后执行。"),
        ("标记大小", "调整小十字的可见尺寸", "默认 0.005 m；不要把它当穴位范围。"),
        ("显示姿态滑动条", "打开项目姿态控制", "SMPL-X 可用；SKEL 用专属控制面板。"),
    ],
    [2100, 3100, 4160],
)
add_callout(doc, "元数据修改", "在列表中选中点后，直接修改下方字段；字段改动会自动写入当前 `.blend` 场景，无需另按“应用”。修改后按 Ctrl+S，并观察人体旁名称是否同步。", "info")
doc.add_heading("7.2 `.blend` 与 `.json` 的区别", 2)
add_table(
    doc,
    ["文件", "保存什么", "用途"],
    [
        ("`.blend`", "完整场景、人体、体型、姿态、点和界面状态", "继续编辑、复核、留档。"),
        ("`.json`", "模型签名、体型/姿态参数、点的表面绑定、名称、编码、侧别、备注与坐标", "程序处理、版本比较、数据交接。"),
    ],
    [1600, 4760, 3000],
)
add_callout(doc, "必须成对保存", "只保存 JSON 不足以完整恢复医生当时的可视环境；只保存 `.blend` 又不利于数据处理。每个版本都应同时保存两者。", "warning")

page_break(doc)
doc.add_heading("8. 保存、导出与重新打开复查", 1)
doc.add_heading("8.1 每完成一个区域", 2)
add_bullets(doc, ["按 Ctrl+S 保存工作副本。", "滚动到“4. 数据交换”，单击“导出穴位 JSON”。", "JSON 使用与 `.blend` 相同主文件名；不要覆盖上一版本。", "在文件管理器中确认两个文件都已生成且不是 0 KB。"], numbered=True)
add_callout(doc, "JSON 文件窗口", "单击导出后会出现 Blender 文件浏览器。先进入与工作副本相同的输出目录；在底部文件名框输入同名主文件名，保留 `.json` 扩展名；单击右下角确认导出。若弹出覆盖提示，只有在明确更新同一版本时才确认，否则取消并递增版本号。", "info")
add_bullets(doc, ["成功后 Blender 底部状态栏会短暂显示导出完成/文件路径。", "按 Windows 键 + E 打开文件资源管理器，进入刚才目录；确认 `.blend` 与 `.json` 同名且大小大于 0 KB。", "建议的本机输出目录：在第 1.1 节写明的项目根目录下新建 `医生标注输出/标注者编号/日期/`；若课题另有数据目录，以负责人书面通知为准。"])
doc.add_heading("8.2 关机前验收", 2)
add_bullets(doc, ["关闭 Blender；如果弹出未保存提示，先保存。", "重新双击刚才的工作副本，而不是空白模板。", "按 Home、N，打开“全身穴位标注”，确认点数。", "用“定位”抽查至少 3 点：一个中线、一个左侧、一个右侧；包含手/面时再抽查一个精细点。", "必要时新开空白同类模板，导入 JSON；只有拓扑相同的模型才能导入。"], numbered=True)
add_callout(doc, "`.blend` 双击没有打开", "先启动项目的一键 `.cmd`，在 Blender 中选择“文件 → 打开”，找到自己的工作副本。出现“是否保存当前文件”时，不要覆盖刚启动的空白模板；选择不保存后再打开工作副本。", "warning")

page_break(doc)
doc.add_heading("9. SKEL：用途、外观与显示原则", 1)
add_figure(doc, "05_skel_open_and_annotate.jpg", "图 7　SKEL 默认只显示皮肤。穴位标注面板与 SMPL-X 共用。")
add_para(doc, "SKEL 的女性/男性外观来自你选择的授权模型模板。它不是软件对患者性别、体型或疾病的判断；当前模板只是研究坐标与生物力学参考。")
add_callout(doc, "何时用 SKEL", "需要把表面点与骨骼、关节或生物力学结构做对照时使用。只做常规全身、手部或头面部落点时，SMPL-X 更直接。", "info")
add_callout(doc, "SKEL 没有中性模板", "研究方案未规定 SKEL 性别时，医生不要自行选择。由研究负责人先在任务单中明确女性或男性模板；两者不能被当成同一个参考坐标。", "warning")
add_bullets(doc, ["医生落点时默认仅显示皮肤。", "需要解剖对照时临时显示骨骼；完成对照后再次隐藏。", "关节名称默认关闭，因为会遮挡皮肤和穴位名称。", "SKEL 没有独立手指关节与面部表情，相关区域改用 SMPL-X。"])

page_break(doc)
doc.add_heading("10. SKEL：调整体型、姿态并更新", 1)
add_figure(doc, "06_skel_shape_visibility.jpg", "图 8　SKEL 模型显示与 β0～β9 体型参数。")
doc.add_heading("10.1 体型参数", 2)
add_bullets(doc, ["β0 约对应身高方向，β1 约对应体重方向；其余是统计体型分量。", "先从 0 小幅改变，建议常规试验范围限制在 -2～2。", "每次只改一个参数，更新后观察正、侧、背三面；确认效果再改下一个。", "参数并非厘米、千克或患者测量值，必须在研究记录中保存具体 β 数值。"])
add_figure(doc, "07_skel_pose_update.jpg", "图 9　SKEL 常用姿态滑条与更新按钮。改变参数后必须计算刷新。")
doc.add_heading("10.2 姿态与计算", 2)
add_bullets(doc, ["只使用“SKEL 生物力学控制”中的姿态滑条；“全身穴位标注”中的 SMPL 骨骼滑条不适用于 SKEL，当前版本会自动隐藏。", "滑条单位为度；拖动时只修改参数，不会立即驱动人体。先调腰、胸、颈，再调肩肘等局部。", "只看皮肤时单击“快速更新皮肤”；显示骨骼或需要关节对照时单击“完整更新皮肤和骨骼”。", "更新完成、模型不再变化后再落点。", "若改了参数但忘记更新，屏幕上的皮肤仍是旧状态；此时不要标注。"])
add_callout(doc, "SKEL 试调与等待", "一次只改一个 β 或姿态滑条，先小幅正值/负值试探，方向不符就回到原值。更新期间不要继续拖动或落点；按钮触发计算后等待界面恢复响应、人体停止变化，并看到信息提示后再继续。完整更新会明显慢于快速更新。", "warning")
page_break(doc)
doc.add_heading("11. SKEL：标注、保存与复核", 1)
add_para(doc, "SKEL 与 SMPL-X 使用同一个“全身穴位标注”面板，所以填写、选点、列表管理、保存 `.blend` 和导出 JSON 的步骤完全相同。差别只在体型/姿态的计算更新与骨骼显示。")
add_table(
    doc,
    ["顺序", "SKEL 必做动作"],
    [
        ("1", "另存为工作副本。"),
        ("2", "在“SKEL 生物力学控制”中调整显示、β 参数和姿态。"),
        ("3", "执行快速更新或完整更新，等待完成。"),
        ("4", "回到“全身穴位标注”，填写字段并落点。"),
        ("5", "如再次改变 β/姿态，重新更新，并刷新、复核所有已有点。"),
        ("6", "保存 `.blend`、导出 JSON、关闭后重开抽查。"),
    ],
    [900, 8460],
)
add_callout(doc, "两个重置按钮", "“姿态归零并更新”只清零姿态并立即重算皮肤，保留体型 β；“恢复默认体型和姿势”会同时清除 β 与姿态并重算皮肤。使用前先 Ctrl+S 保存一个可回退版本。", "warning")

page_break(doc)
doc.add_heading("12. 医生标注质量控制", 1)
add_table(
    doc,
    ["检查项", "通过标准"],
    [
        ("模型", "已记录 SMPL-X/SKEL、性别模板和版本；没有混用拓扑。"),
        ("体型与姿态", "参数已更新；记录了关键身高体重或 β/姿态值。"),
        ("名称与编码", "每个点都有统一编码和医生认可的中文名称。"),
        ("身体分区/区域", "分区、经脉或研究区域填写一致。"),
        ("侧别", "左、右、中线按人体自身方向判断；双侧点分开落点。"),
        ("表面位置", "至少两个视角复核，没有穿到背面或落到另一肢体。"),
        ("手部/头面", "使用 SMPL-X，已局部放大并做正侧面检查。"),
        ("保存", "`.blend` 与 `.json` 同名成对存在；重开后点数一致。"),
    ],
    [3000, 6360],
)
doc.add_heading("12.1 推荐双人流程", 2)
add_bullets(doc, ["医生 A 独立落点并导出 v01。", "医生 B 只读复核名称、编码、侧别与位置；把分歧写入备注或复核表。", "达成一致后由指定人员修改，另存 v02；保留 v01，不覆盖原始判断。", "研究负责人确认后锁定版本。"])

page_break(doc)
doc.add_heading("13. 常见问题速查", 1)
add_callout(doc, "拓扑不兼容", "SMPL、SMPL-X、SKEL 以及 SMPL-X 的不同拓扑数据不能混用。若导入提示模型签名不一致，请打开与导出时完全相同的模型模板。", "warning")
add_table(
    doc,
    ["问题", "原因", "处理"],
    [
        ("点显示成很大球体", "旧版模板/旧插件，或标记大小过大", "使用当前模板；在显示设置把标记大小设为 0.005 m。"),
        ("显示 ACU_随机字符", "旧版对象名，没有填写医学名称", "在已标注列表选中，补填中文名称；必要时重建该点。"),
        ("人体太小", "视图没有聚焦", "选中人体，在视图区按 Home。"),
        ("没有右侧栏", "N 面板关闭", "鼠标放在视图区按 N。"),
        ("没有穴位面板", "不在 Item/项目标签", "点击右侧竖排“项目/Item”。"),
        ("找不到姿态滑动条", "停留在“SMPL模型”标签，或尚未展开控件", "切到右侧竖排“项目/Item”，在“全身穴位标注”中向下滚动并单击“显示姿态滑动条”。"),
        ("SMPL-X 出现两个人体", "误按了“添加”", "关闭且不保存，重新打开工作副本；模板内无需添加。"),
        ("点落错了", "表面选择错误", "删除该点，在正确表面重新标；不能直接拖动。"),
        ("改姿态后点偏了", "没有刷新或未复核", "更新模型/姿态修正，按“刷新”，逐点复核。"),
        ("SKEL 改参数没变化", "尚未执行计算", "快速更新皮肤，或完整更新皮肤和骨骼。"),
        ("SKEL 字太多遮挡", "骨骼或关节名称开启", "关闭骨骼/关节名称并应用显示设置。"),
        ("JSON 导不进去", "模型拓扑签名不同", "打开导出时同一模型族、同一拓扑的模板。"),
        ("找不到上次点位", "打开了空白模板", "打开自己的 `.blend` 工作副本；必要时导入配套 JSON。"),
    ],
    [2400, 2700, 4260],
)

page_break(doc)
doc.add_heading("14. 数据边界与合规提醒", 1)
add_bullets(
    doc,
    [
        "本工具用于科研标注，不替代医生诊断，不自动推荐穴位，也不验证医学正确性。",
        "当前点位是模型皮肤表面上的研究标注；Blender 世界坐标不是机器人可直接执行的坐标。机器人使用前仍需患者配准、坐标转换、安全距离、压力/力控和临床验证。",
        "模型体型是参数化统计近似，不等同患者扫描、触诊结果或组织深度。",
        "SMPL-X、SKEL 及其数据受各自许可约束。把工具或模型复制到其他电脑、机构或外部人员前，应先确认对方有相应授权。",
        "避免在文件名和普通备注中写入不必要的患者身份信息；按课题的数据管理方案存储。",
    ],
)
add_callout(doc, "当前未包含", "自动穴位推荐、经络曲线、区域涂绘、患者扫描拟合和机器人坐标输出不在本教程与当前插件范围内。", "warning")

page_break(doc)
doc.add_heading("附录 A：交付前一分钟检查表", 1)
add_table(
    doc,
    ["勾选", "检查内容"],
    checkbox_rows(
        [
            "我打开的是自己的工作副本，不是空白模板。",
            "模型族、性别模板、体型和姿态已记录。",
            "每个点都有标准编码、中文名称、分区、区域与侧别。",
            "左右方向按人体自身方向确认。",
            "每个点至少从两个视角复核；手部/头面已放大复核。",
            "改体型/姿态后已更新模型并刷新点位。",
            "`.blend` 已 Ctrl+S 保存。",
            "同名 `.json` 已导出，两个文件均非 0 KB。",
            "关闭并重开 `.blend` 后，点数和抽查位置正确。",
            "文件名不含患者姓名；版本号清楚。",
        ]
    ),
    [900, 8460],
)
add_callout(doc, "完成标准", "以上 10 项全部勾选后，才把该版本交给研究负责人或下一位复核医生。", "success")

doc.add_heading("附录 B：本教程实测范围", 1)
add_para(doc, "教程制作时已对 SMPL-X 中性/女性/男性模板、SKEL 女性/男性模板执行自动化流程测试；验证了 3 点创建、名称修改、体型/姿态更新、`.blend` 保存后重开、JSON 导出/导入、重复导入保护和错误拓扑拒绝。")
add_para(doc, "软件环境：Blender 4.5.12 LTS；穴位标注插件 v0.3.2。", color=MUTED, italic=True)

doc.core_properties.title = "医生版 SMPL-X 与 SKEL 穴位标注图文教程"
doc.core_properties.subject = "科研阶段医生人体表面穴位标注操作手册"
doc.core_properties.author = "按摩理疗机器人项目组"
doc.core_properties.keywords = "SMPL-X, SKEL, Blender, 穴位标注, 医生教程"
doc.save(OUTPUT)
print(f"DOCX={OUTPUT}")
